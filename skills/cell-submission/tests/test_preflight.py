import json
import sys
import tempfile
import unittest
from pathlib import Path


SCRIPT_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPT_DIR))

from preflight import GATES, check, file_in_root, scan, sha256  # noqa: E402


class SubmissionPreflightTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source_dir = self.root / "source"
        self.delivery_dir = self.root / "deliverables"
        self.work_dir = self.root / "work"
        self.source_dir.mkdir()
        self.delivery_dir.mkdir()
        self.work_dir.mkdir()

        self.source = self.source_dir / "Manuscript.txt"
        self.source.write_text("Original result: n=24 and P=0.031.\n", encoding="utf-8")
        self.delivery = self.delivery_dir / "Manuscript.txt"
        self.delivery.write_text("Final result: n=24 and P=0.031.\n", encoding="utf-8")
        self.manifest_path = self.work_dir / "package_manifest.json"
        self.write_manifest()

    def tearDown(self):
        self.temp.cleanup()

    def manifest(self):
        checks = {}
        for gate in GATES:
            checks[gate] = {
                "status": "not_applicable" if gate in {"anonymity"} else "pass",
                "evidence": f"Verified {gate} against the final package.",
            }
        return {
            "version": 2,
            "root": "../deliverables",
            "journal": "Example Journal",
            "article_type": "Original Article",
            "stage": "initial_submission",
            "rules_checked_at": "2026-09-28",
            "rules": [{
                "id": "JR-001",
                "topic": "editable manuscript",
                "requirement": "Submit one editable manuscript file.",
                "strength": "required",
                "source_kind": "journal_official",
                "source_locator": "https://example.test/journal/author-guide",
                "source_section": "Initial submission > Files",
                "accessed_at": "2026-09-28",
                "applies_to": {
                    "article_type": "Original Article",
                    "stage": "initial_submission",
                },
                "status": "verified_applied",
                "interpretation": "The final editable manuscript satisfies this file requirement.",
                "target_files": ["Manuscript.txt"],
            }],
            "checks": checks,
            "blockers": [],
            "sources": [{"path": "../source/Manuscript.txt", "sha256": sha256(self.source)}],
            "files": [{
                "path": "Manuscript.txt",
                "basis": "journal_required",
                "rule_ids": ["JR-001"],
                "evidence": "Author guide requires an editable manuscript.",
                "role": "clean_manuscript",
                "clean": True,
                "sha256": sha256(self.delivery),
            }],
        }

    def write_manifest(self, manifest=None):
        self.manifest_path.write_text(
            json.dumps(manifest or self.manifest(), ensure_ascii=False), encoding="utf-8"
        )

    def test_valid_minimal_package_passes(self):
        result = check(self.manifest_path)
        self.assertEqual(result["status"], "PASS")
        self.assertEqual(result["errors"], [])
        self.assertIn("Manuscript.txt", result["files_checked"])

    def test_placeholder_blocks_package(self):
        self.delivery.write_text("[[TODO: add ethics approval]]\n", encoding="utf-8")
        manifest = self.manifest()
        manifest["files"][0]["sha256"] = sha256(self.delivery)
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertEqual(result["status"], "BLOCKED")
        self.assertTrue(any("Unresolved placeholder" in item for item in result["errors"]))

    def test_changed_source_and_stale_final_hash_are_blocked(self):
        self.source.write_text("Source changed after freezing.\n", encoding="utf-8")
        self.delivery.write_text("Final changed after review.\n", encoding="utf-8")
        result = check(self.manifest_path)
        self.assertTrue(any("Source is missing or changed" in item for item in result["errors"]))
        self.assertTrue(any("Invalid/stale final hash" in item for item in result["errors"]))

    def test_extra_unlisted_file_is_blocked(self):
        (self.delivery_dir / "internal_notes.txt").write_text("not for delivery", encoding="utf-8")
        result = check(self.manifest_path)
        self.assertTrue(any("Extra file outside whitelist" in item for item in result["errors"]))

    def test_rule_ledger_is_required(self):
        manifest = self.manifest()
        manifest["rules"] = []
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertTrue(any("scoped journal rule" in item for item in result["errors"]))

    def test_rule_must_match_article_type_and_stage(self):
        manifest = self.manifest()
        manifest["rules"][0]["applies_to"]["stage"] = "revision"
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertTrue(any("submission stage" in item for item in result["errors"]))

    def test_required_rule_cannot_be_marked_not_applicable(self):
        manifest = self.manifest()
        manifest["rules"][0]["status"] = "verified_not_applicable"
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertTrue(any("required rule" in item for item in result["errors"]))

    def test_file_must_link_to_matching_rule(self):
        manifest = self.manifest()
        manifest["files"][0]["rule_ids"] = ["JR-999"]
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertTrue(any("Unknown rule ID" in item for item in result["errors"]))
        self.assertTrue(any("basis is not supported" in item for item in result["errors"]))

    def test_rule_target_must_be_delivered(self):
        manifest = self.manifest()
        manifest["rules"][0]["target_files"] = ["Missing.docx"]
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertTrue(any("unlisted file" in item for item in result["errors"]))

    def test_rule_dates_cannot_be_future_or_after_review(self):
        manifest = self.manifest()
        manifest["rules"][0]["accessed_at"] = "2099-01-01"
        self.write_manifest(manifest)
        result = check(self.manifest_path)
        self.assertTrue(any("future" in item for item in result["errors"]))

    def test_legacy_manifest_is_rejected(self):
        manifest = self.manifest()
        manifest["version"] = 1
        self.write_manifest(manifest)
        with self.assertRaisesRegex(ValueError, "version must be 2"):
            check(self.manifest_path)

    def test_unsafe_deliverable_path_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "Unsafe deliverable path"):
            file_in_root(self.delivery_dir.resolve(), "../source/Manuscript.txt")

    def test_text_scan_reports_numbers_and_review_markers(self):
        path = self.root / "draft.md"
        path.write_text("Participants: 18.\nTODO\nEffect was 12.5%.\n", encoding="utf-8")
        result = scan(path)
        self.assertEqual([item["value"] for item in result["numeric_contexts"]], ["18", "12.5%"])
        self.assertIn("TODO", result["review_terms"])
        self.assertTrue(result["placeholders"])

    def test_real_docx_is_scanned(self):
        from docx import Document

        path = self.root / "manuscript.docx"
        document = Document()
        document.core_properties.author = "Example Author"
        document.add_paragraph("A measured value of 7.2 was observed.")
        document.save(path)

        result = scan(path)
        self.assertEqual(result["numeric_contexts"][0]["value"], "7.2")
        self.assertTrue(any("Example Author" == value for value in result["metadata"].values()))


if __name__ == "__main__":
    unittest.main()
