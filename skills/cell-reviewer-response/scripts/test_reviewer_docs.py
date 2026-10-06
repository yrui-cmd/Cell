#!/usr/bin/env python3
"""Synthetic implementation tests. Never deliver generated fixtures to authors."""
from __future__ import annotations

import copy
import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from reviewer_docs import (CHECKS, ValidationError, build, sha256, validate,
                           verification_fingerprint)


def fixture(root: Path) -> dict:
    root.mkdir(parents=True, exist_ok=True)
    manuscript = root / "manuscript.txt"
    manuscript.write_text(
        "SYNTHETIC TEST MANUSCRIPT, NOT RESEARCH DATA\n"
        "Methods: Each point represents one independent specimen.\n"
        "Results: Group-specific distributions are reported in Figure 2.\n"
        "Discussion: The study design does not establish causality.\n"
        "Data availability: Data are included in the supplementary material.\n",
        encoding="utf-8")
    texts = [
        "Please clarify the data availability statement.",
        "In Figure 2, what does each point represent?",
        "Please report the group-specific distributions in Figure 2 and clarify whether Figure S1 establishes causality.",
        "In Figure 2, explain the unit of replication.",
    ]
    review = root / "reviews.txt"
    review.write_text("SYNTHETIC TEST REVIEWS\n\n" + "\n\n".join(texts), encoding="utf-8")
    ref = {"artifact_id": "manuscript", "locator": "Methods, paragraph 1",
           "supports": "The unit represented by each point is described.",
           "excerpt": "Each point represents one independent specimen."}
    data = {
        "schema_version": 1,
        "meta": {"title": "Synthetic document test", "manuscript_id": "TEST-ONLY",
                 "original_comment_count": 4, "review_sources_complete": True,
                 "current_manuscript_id": "manuscript",
                 "completion_status": "needs_author_input",
                 "opening": "Thank you for reviewing the manuscript. We respond to each comment below."},
        "artifacts": [
            {"id": "reviews", "path": str(review), "sha256": sha256(review)},
            {"id": "manuscript", "path": str(manuscript), "sha256": sha256(manuscript)},
        ],
        "reviewers": [
            {"id": "E", "label": "Editor", "kind": "editor", "order": 1, "comment_count": 1},
            {"id": "R1", "label": "Reviewer 1", "kind": "reviewer", "order": 2, "comment_count": 2},
            {"id": "R2", "label": "Reviewer 2", "kind": "reviewer", "order": 3, "comment_count": 1},
        ],
        "comments": [], "issues": [], "checks": {},
    }
    mappings = [("E-C1", "E", 1), ("R1-C1", "R1", 1), ("R1-C2", "R1", 2), ("R2-C1", "R2", 1)]
    for (cid, rid, order), original in zip(mappings, texts):
        data["comments"].append({"id": cid, "reviewer_id": rid, "order": order,
                                 "original_label": f"Comment {order}", "original_text": original,
                                 "source_artifact_id": "reviews", "source_locator": f"{rid}, comment {order}"})
    specs = [
        ("E-C1a", "E-C1", "General / Data availability", False, "核对数据可得性说明。"),
        ("R1-C1a", "R1-C1", "Figure 2", False, "解释每个数据点所代表的单位。"),
        ("R1-C2a", "R1-C2", "Figure 2", True, "补充不同组别的完整分布信息。"),
        ("R1-C2b", "R1-C2", "Supplementary Figure S1", True, "界定补充图能否支持因果结论。"),
        ("R2-C1a", "R2-C1", "Figure 2", False, "区分独立样本与重复测量。"),
    ]
    for iid, cid, group, pending, question in specs:
        item = {"id": iid, "comment_id": cid, "primary_group": group,
                "panel_refs": [group] if "Figure" in group else [],
                "related_groups": ["General / Statistics"] if iid == "R1-C2a" else [],
                "question": question, "route": "reanalysis" if pending else "reply_only",
                "category": ("analysis" if pending else "clarification"),
                "severity": ("high" if pending else "medium"),
                "status": "open" if pending else "ready",
                "draft_response": ("The additional analysis remains to be assessed before the response is finalized."
                                   if pending else "Each point represents one independent specimen, as described in the Methods."),
                "evidence": [] if pending else [copy.deepcopy(ref)]}
        if pending:
            item["plan"] = {
                "action": "核对相关结果并据实更新本条回复。",
                "inputs": "与当前问题相对应的实际结果及最终修订稿。",
                "design": "对照原始要求核实比较对象、统计单位及结论范围。",
                "output": "更新相关分图说明或 Discussion 中的对应表述。",
                "acceptance": "所提供证据能直接回答该子问题，且定位与最终稿一致。",
                "fallback": "若证据不支持原解释，收窄结论并明确相应局限。"}
        if cid == "E-C1":
            excerpt = "Data are included in the supplementary material."
            item["draft_response"] = "The data availability statement specifies that the data are included in the supplementary material."
            item["evidence"] = [{"artifact_id": "manuscript", "locator": "Data availability statement",
                                 "supports": excerpt, "excerpt": excerpt}]
        data["issues"].append(item)
    return data


def completed(data: dict) -> dict:
    data = copy.deepcopy(data)
    data["meta"]["completion_status"] = "ready_to_submit"
    data["meta"]["principal_revisions"] = [
        "Clarified the experimental unit used in Figure 2.",
        "Reported the group-specific distributions and limited the causal interpretation.",
    ]
    for issue in data["issues"]:
        if issue["status"] == "open":
            causal = issue["id"].endswith("b")
            locator = "Discussion, paragraph 1" if causal else "Results, paragraph 1; Figure 2"
            excerpt = ("The study design does not establish causality." if causal
                       else "Group-specific distributions are reported in Figure 2.")
            ref = {"artifact_id": "manuscript", "locator": locator,
                   "supports": excerpt, "excerpt": excerpt}
            issue["status"] = "resolved"
            issue["resolution"] = {"disposition": "implemented", "summary": excerpt,
                                   "evidence": [copy.deepcopy(ref)], "changes": [copy.deepcopy(ref)]}
    for c in data["comments"]:
        kids = [i for i in data["issues"] if i["comment_id"] == c["id"]]
        c["final_issue_ids"] = [i["id"] for i in kids]
        c["response_mode"] = "agree_and_change" if any(i["status"] == "resolved" for i in kids) else "clarification"
        if c["id"] == "E-C1":
            c["final_response"] = "The data availability statement specifies that the data are included in the supplementary material."
        elif c["id"] == "R1-C2":
            c["final_response"] = ("Group-specific distributions are reported in Figure 2. "
                                   "The Discussion clarifies that the study design does not establish causality.")
        else:
            c["final_response"] = "Each point represents one independent specimen, as described in the Methods."
    method_text = "Each point represents one independent specimen."
    data["comments"][1]["revised_text"] = method_text
    data["comments"][1]["revised_text_evidence"] = {
        "artifact_id": "manuscript", "locator": "Methods, paragraph 1",
        "supports": "The quoted manuscript text states the experimental unit.",
        "excerpt": method_text,
    }
    data["checks"] = {key: True for key in CHECKS}
    seal(data)
    return data


def seal(data: dict) -> None:
    data["verification"] = {
        "checked_at": "2026-09-28",
        "content_sha256": verification_fingerprint(data),
    }


class ReviewerDocsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.data = fixture(self.root)

    def tearDown(self):
        self.tmp.cleanup()

    def test_pending_blocks_final_file(self):
        target = self.root / "blocked.docx"
        with self.assertRaises(ValidationError):
            build(self.data, self.root, "stage2", target)
        self.assertFalse(target.exists())

    def test_stage1_red_and_primary_counts(self):
        target = self.root / "first.docx"
        result = build(self.data, self.root, "stage1", target)
        self.assertEqual((result["comments"], result["issues"]), (4, 5))
        doc = Document(target)
        red = [r.text for p in doc.paragraphs for r in p.runs if str(r.font.color.rgb) == "C00000"]
        self.assertTrue(any("R1-C2a" in value for value in red))
        self.assertTrue(any("解决动作" in value for value in red))
        rows = [[c.text for c in r.cells] for r in doc.tables[0].rows]
        self.assertIn(["Figure 2", "3", "2", "1"], rows)
        self.assertFalse(any(row[0] == "General / Statistics" for row in rows))

    def test_ready_final_order_and_no_red(self):
        target = self.root / "final.docx"
        build(completed(self.data), self.root, "stage2", target)
        doc = Document(target)
        texts = [p.text for p in doc.paragraphs]
        self.assertLess(texts.index("Editor"), texts.index("Reviewer 1"))
        self.assertLess(texts.index("Reviewer 1"), texts.index("Reviewer 2"))
        self.assertFalse(any(str(r.font.color.rgb) == "C00000" for p in doc.paragraphs for r in p.runs))
        self.assertIn("Summary of principal revisions", texts)
        self.assertIn("Comment", texts)
        self.assertIn("Changes made", texts)
        self.assertIn("Revised manuscript text", texts)
        self.assertIn("Locations", texts)
        self.assertFalse(any("R1-C2a" in s for s in texts))

    def test_missing_subquestion_fails(self):
        data = completed(self.data)
        data["comments"][2]["final_issue_ids"].pop()
        with self.assertRaises(ValidationError):
            validate(data, self.root, "stage2")

    def test_source_version_drift_fails(self):
        Path(self.data["artifacts"][1]["path"]).write_text("changed", encoding="utf-8")
        with self.assertRaises(ValidationError):
            validate(self.data, self.root, "stage1")

    def test_changed_final_response_invalidates_checks(self):
        data = completed(self.data)
        data["comments"][0]["final_response"] += " Changed after review."
        with self.assertRaisesRegex(ValidationError, "verification is stale"):
            validate(data, self.root, "stage2")

    def test_docx_equation_requires_bound_object_review(self):
        data = completed(self.data)
        path = self.root / "manuscript.docx"
        doc = Document()
        doc.add_paragraph(Path(data["artifacts"][1]["path"]).read_text(encoding="utf-8"))
        equation = doc.paragraphs[0]._p.makeelement(
            "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath")
        doc.paragraphs[0]._p.append(equation)
        doc.save(path)
        artifact = data["artifacts"][1]
        artifact.update(path=str(path), sha256=sha256(path))
        seal(data)
        with self.assertRaisesRegex(ValidationError, "object_review"):
            validate(data, self.root, "stage2")
        artifact["object_review"] = {
            "source_sha256": artifact["sha256"],
            "inspected_types": ["equation"],
            "notes": "Inspected the equation against the final rendered manuscript.",
        }
        validate(data, self.root, "stage2")

    def test_original_comment_cannot_be_rewritten(self):
        self.data["comments"][1]["original_text"] = "A paraphrase that was not in the review."
        with self.assertRaises(ValidationError):
            validate(self.data, self.root, "stage1")

    def test_unverified_humanizer_fails(self):
        data = completed(self.data)
        data["checks"]["humanizer"] = False
        with self.assertRaises(ValidationError):
            validate(data, self.root, "stage2")

    def test_no_locations_fails(self):
        data = completed(self.data)
        data["issues"][2]["resolution"]["changes"] = []
        with self.assertRaises(ValidationError):
            validate(data, self.root, "stage2")

    def test_incomplete_sources_only_stage1(self):
        data = completed(self.data)
        data["meta"]["review_sources_complete"] = False
        data["meta"]["source_gaps"] = "Reviewer 3 的附件尚未提供。"
        validate(data, self.root, "stage1")
        with self.assertRaises(ValidationError):
            validate(data, self.root, "stage2")

    def test_reasoned_disagreement_allowed(self):
        data = completed(self.data)
        issue = data["issues"][3]
        issue["route"] = "reasoned_disagreement"
        issue["resolution"]["disposition"] = "reasoned_disagreement"
        issue["resolution"]["changes"] = []
        issue["resolution"]["no_change_reason"] = "The manuscript already makes the scope of inference explicit."
        seal(data)
        validate(data, self.root, "stage2")

    def test_placeholder_blocks_final(self):
        data = completed(self.data)
        data["comments"][0]["final_response"] = "We have made this change on page XX."
        with self.assertRaises(ValidationError):
            validate(data, self.root, "stage2")

    def test_placeholder_in_change_location_fails(self):
        data = completed(self.data)
        data["issues"][2]["resolution"]["changes"][0]["locator"] = "Page XX, lines XX"
        with self.assertRaises(ValidationError):
            validate(data, self.root, "stage2")

    def test_placeholder_in_change_summary_fails(self):
        data = completed(self.data)
        data["issues"][2]["resolution"]["summary"] = "Analysis complete; add result TODO."
        with self.assertRaisesRegex(ValidationError, "final resolution contains a placeholder"):
            validate(data, self.root, "stage2")

    def test_duplicate_issue_fails(self):
        self.data["issues"].append(copy.deepcopy(self.data["issues"][0]))
        with self.assertRaises(ValidationError):
            validate(self.data, self.root, "stage1")

    def test_input_is_never_overwritten(self):
        with self.assertRaises(ValidationError):
            build(self.data, self.root, "stage1", Path(self.data["artifacts"][1]["path"]))

    def test_parent_count_mismatch_fails(self):
        self.data["meta"]["original_comment_count"] = 5
        with self.assertRaises(ValidationError):
            validate(self.data, self.root, "stage1")

    def test_ready_status_cannot_hide_open_issue(self):
        self.data["meta"]["completion_status"] = "ready_to_submit"
        with self.assertRaisesRegex(ValidationError, "incompatible with unresolved issues"):
            validate(self.data, self.root, "stage1")

    def test_final_response_mode_required(self):
        data = completed(self.data)
        data["comments"][0].pop("response_mode")
        with self.assertRaisesRegex(ValidationError, "response_mode"):
            validate(data, self.root, "stage2")

    def test_revised_text_must_match_current_manuscript(self):
        data = completed(self.data)
        data["comments"][1]["revised_text"] = "A sentence that is not in the manuscript."
        with self.assertRaisesRegex(ValidationError, "exactly match"):
            validate(data, self.root, "stage2")

    def test_category_and_severity_required(self):
        self.data["issues"][0].pop("category")
        self.data["issues"][1]["severity"] = "urgent"
        with self.assertRaisesRegex(ValidationError, "category"):
            validate(self.data, self.root, "stage1")

    def test_stage2_requires_principal_revision_summary(self):
        data = completed(self.data)
        data["meta"]["principal_revisions"] = []
        with self.assertRaisesRegex(ValidationError, "principal_revisions"):
            validate(data, self.root, "stage2")

    def test_unable_mode_requires_verified_alternative(self):
        data = completed(self.data)
        data["comments"][0]["response_mode"] = "unable_with_alternative"
        with self.assertRaisesRegex(ValidationError, "verified alternative resolution"):
            validate(data, self.root, "stage2")

    def test_verified_limited_alternative_is_allowed(self):
        data = completed(self.data)
        issue = data["issues"][0]
        issue["route"] = "alternative_resolution"
        issue["status"] = "resolved"
        issue["resolution"] = {
            "disposition": "limited_alternative",
            "summary": "The existing data-availability statement is retained as the supported alternative.",
            "evidence": copy.deepcopy(issue["evidence"]),
            "changes": [],
            "no_change_reason": "The current statement already reports the supported availability boundary.",
        }
        data["comments"][0]["response_mode"] = "unable_with_alternative"
        seal(data)
        validate(data, self.root, "stage2")


if __name__ == "__main__":
    unittest.main(verbosity=2)
