#!/usr/bin/env python3
"""Build two-stage reviewer-response DOCX files from an agent-verified ledger.

No model/API calls. This module checks structure, source snapshots and gates.
It does not infer whether an experiment was scientifically adequate.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

RED = "C00000"
BLACK = "000000"
ROUTES = {"reply_only", "manuscript_edit", "figure_fix", "reanalysis",
          "new_experiment", "evidence_lookup", "reasoned_disagreement"}
CHECKS = ("coverage", "evidence", "locations", "figures",
          "humanizer", "facts_after_humanizer")
PLACEHOLDER = re.compile(
    r"\b(?:TBD|TODO|PLACEHOLDER)\b|\[(?:INSERT|待补|待填|待确认)[^\]]*\]"
    r"|\b(?:page|pages|line|lines)\s+(?:XX+|\?+)\b|待取得.{0,24}补写"
    r"|(?:cite|filecite)", re.I)


class ValidationError(ValueError):
    pass


def text(value: Any) -> str:
    return value.strip() if isinstance(value, str) else ""


def norm(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def sha256(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest() if hasattr(hashlib, "file_digest") else _old_hash(stream)


def _old_hash(stream: Any) -> str:
    digest = hashlib.sha256()
    for chunk in iter(lambda: stream.read(1024 * 1024), b""):
        digest.update(chunk)
    return digest.hexdigest()


def artifact_text(path: Path) -> str | None:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".tsv", ".json"}:
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".docx":
        doc = Document(path)
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells)
        return "\n".join(chunks)
    return None  # Figures/PDFs must be read by the host; never silently OCR.


def resolve_path(raw: str, base: Path) -> Path:
    path = Path(raw).expanduser()
    return (path if path.is_absolute() else base / path).resolve()


def natural(value: str) -> tuple[Any, ...]:
    lower = value.lower()
    rank = (0 if lower.startswith("figure ") else
            1 if lower.startswith("supplementary figure ") else
            2 if lower.startswith("table ") else
            3 if lower.startswith("supplementary table ") else 4)
    parts = tuple((0, int(s)) if s.isdigit() else (1, s)
                  for s in re.split(r"(\d+)", lower))
    return rank, parts


def validate(data: dict[str, Any], base: Path, stage: str) -> dict[str, Any]:
    errors: list[str] = []
    artifacts: dict[str, dict[str, Any]] = {}
    source_texts: dict[str, str | None] = {}
    if data.get("schema_version") != 1:
        errors.append("schema_version must be 1")
    meta = data.get("meta", {})
    for a in data.get("artifacts", []):
        aid = text(a.get("id"))
        if not aid or aid in artifacts:
            errors.append("artifact IDs must be nonempty and unique")
            continue
        artifacts[aid] = a
        raw = text(a.get("path"))
        path = resolve_path(raw, base) if raw else base
        if not raw or not path.is_file():
            errors.append(f"{aid}: source file missing")
            continue
        if not re.fullmatch(r"[0-9a-f]{64}", str(a.get("sha256", ""))):
            errors.append(f"{aid}: SHA-256 snapshot missing")
        elif sha256(path) != a["sha256"]:
            errors.append(f"{aid}: source version changed; re-read and reverify")
        try:
            source_texts[aid] = artifact_text(path)
        except Exception as exc:
            errors.append(f"{aid}: source unreadable ({exc})")

    def reference(ref: Any, label: str) -> None:
        if not isinstance(ref, dict):
            errors.append(f"{label}: evidence must be an object")
            return
        aid = ref.get("artifact_id")
        if aid not in artifacts:
            errors.append(f"{label}: unknown evidence artifact {aid!r}")
        if not text(ref.get("locator")) or not text(ref.get("supports")):
            errors.append(f"{label}: evidence needs locator and supports")
        if stage == "stage2" and PLACEHOLDER.search(text(ref.get("locator"))):
            errors.append(f"{label}: unresolved location placeholder")
        excerpt = text(ref.get("excerpt"))
        content = source_texts.get(aid)
        if excerpt and content is not None and norm(excerpt) not in norm(content):
            errors.append(f"{label}: quoted evidence not found in source")

    reviewers: dict[str, dict[str, Any]] = {}
    reviewer_positions = set()
    for reviewer in data.get("reviewers", []):
        rid = text(reviewer.get("id"))
        position = reviewer.get("order")
        if not rid or rid in reviewers or not text(reviewer.get("label")):
            errors.append("reviewer IDs/labels must be nonempty and IDs unique")
        if not isinstance(position, int) or position in reviewer_positions:
            errors.append("reviewer order must be an explicit unique integer")
        reviewer_positions.add(position)
        reviewers[rid] = reviewer
        if reviewer.get("kind") not in {"editor", "reviewer"}:
            errors.append(f"{rid}: reviewer kind must be editor or reviewer")

    comments: dict[str, dict[str, Any]] = {}
    comment_positions = set()
    for c in data.get("comments", []):
        cid = text(c.get("id"))
        if not cid or cid in comments:
            errors.append("comment IDs must be nonempty and unique")
        comments[cid] = c
        if c.get("reviewer_id") not in reviewers:
            errors.append(f"{cid}: missing reviewer")
        pos = (c.get("reviewer_id"), c.get("order"))
        if not isinstance(c.get("order"), int) or pos in comment_positions:
            errors.append(f"{cid}: duplicate/missing original order")
        comment_positions.add(pos)
        original = text(c.get("original_text"))
        if not original:
            errors.append(f"{cid}: original comment missing")
        aid = c.get("source_artifact_id")
        if aid not in artifacts or not text(c.get("source_locator")):
            errors.append(f"{cid}: original source/locator missing")
        content = source_texts.get(aid)
        if content is None:
            errors.append(f"{cid}: use a readable original text/DOCX or a source-linked transcript")
        elif original and norm(original) not in norm(content):
            errors.append(f"{cid}: original comment does not match source verbatim")
    if not comments:
        errors.append("no reviewer comments supplied")
    expected = meta.get("original_comment_count")
    if not isinstance(expected, int) or expected != len(comments):
        errors.append("original_comment_count does not match extracted parent comments")
    for rid, r in reviewers.items():
        actual = sum(c.get("reviewer_id") == rid for c in comments.values())
        if r.get("comment_count") != actual:
            errors.append(f"{rid}: original comment count mismatch")
    if meta.get("review_sources_complete") is not True and not text(meta.get("source_gaps")):
        errors.append("incomplete review sources require a concrete source_gaps note")

    issue_ids: set[str] = set()
    children: dict[str, list[dict[str, Any]]] = {cid: [] for cid in comments}
    open_ids: list[str] = []
    for issue in data.get("issues", []):
        iid = text(issue.get("id"))
        if not iid or iid in issue_ids:
            errors.append("issue IDs must be nonempty and unique")
        issue_ids.add(iid)
        cid = issue.get("comment_id")
        if cid not in comments:
            errors.append(f"{iid}: orphan issue")
        else:
            children[cid].append(issue)
        for key in ("primary_group", "question", "draft_response"):
            if not text(issue.get(key)):
                errors.append(f"{iid}: {key} missing")
        route, status = issue.get("route"), issue.get("status")
        if route not in ROUTES:
            errors.append(f"{iid}: invalid route")
        if status not in {"ready", "open", "resolved"}:
            errors.append(f"{iid}: invalid status")
        evidence = issue.get("evidence", [])
        for ref in evidence:
            reference(ref, iid)
        if status == "ready":
            if route != "reply_only" or not evidence:
                errors.append(f"{iid}: ready requires reply_only and checked evidence")
        elif status == "open":
            open_ids.append(iid)
            plan = issue.get("plan", {})
            for key in ("action", "inputs", "design", "output", "acceptance", "fallback"):
                if not text(plan.get(key)):
                    errors.append(f"{iid}: actionable plan needs {key}")
        elif status == "resolved":
            resolution = issue.get("resolution", {})
            if not text(resolution.get("summary")) or not resolution.get("evidence"):
                errors.append(f"{iid}: resolution summary/evidence missing")
            if resolution.get("disposition") not in {"implemented", "clarified", "reasoned_disagreement"}:
                errors.append(f"{iid}: resolution disposition missing")
            for ref in resolution.get("evidence", []):
                reference(ref, iid + " resolution")
            changes = resolution.get("changes", [])
            for ref in changes:
                reference(ref, iid + " change")
                if ref.get("artifact_id") != meta.get("current_manuscript_id"):
                    errors.append(f"{iid}: change location must refer to the current manuscript")
            if not changes and not text(resolution.get("no_change_reason")):
                errors.append(f"{iid}: provide actual change locations or a reason no manuscript change is needed")
    for cid, kids in children.items():
        if not kids:
            errors.append(f"{cid}: no issue covers this comment")

    for mapping in data.get("figure_map", []):
        if not text(mapping.get("original")) or not text(mapping.get("revised")):
            errors.append("figure_map needs original and revised labels")
        reference(mapping.get("evidence", {}), "figure_map")

    if stage == "stage2":
        if meta.get("review_sources_complete") is not True:
            errors.append("review sources are incomplete")
        if open_ids:
            errors.append("unresolved issues: " + ", ".join(open_ids))
        for key in CHECKS:
            if data.get("checks", {}).get(key) is not True:
                errors.append(f"final human review check missing: {key}")
        for cid, c in comments.items():
            answer = text(c.get("final_response"))
            if not answer or PLACEHOLDER.search(answer):
                errors.append(f"{cid}: final response is missing or contains a placeholder")
            covered = c.get("final_issue_ids", [])
            expected_ids = {i["id"] for i in children.get(cid, [])}
            if not isinstance(covered, list) or len(covered) != len(set(covered)) or set(covered) != expected_ids:
                errors.append(f"{cid}: final response coverage does not match all subquestions")
        for key in ("title", "manuscript_id", "opening", "signature"):
            if PLACEHOLDER.search(text(meta.get(key))):
                errors.append(f"meta.{key}: final placeholder detected")
        if not text(meta.get("opening")):
            errors.append("a fact-checked short opening is required for stage2")
    if errors:
        raise ValidationError("\n".join(dict.fromkeys(errors)))
    return {"comments": len(comments), "issues": len(issue_ids),
            "open": open_ids, "stage2_ready": stage == "stage2"}


def font(run: Any, red: bool = False, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.color.rgb = RGBColor.from_string(RED if red else BLACK)
    run.bold, run.italic = bold, italic
    props = run._element.get_or_add_rPr()
    fonts = props.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.append(fonts)
    fonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")


def paragraph(doc: Any, value: str, *, red: bool = False, bold: bool = False,
              italic: bool = False, style: str | None = None) -> None:
    for block in value.split("\n"):
        if not block.strip():
            continue
        p = doc.add_paragraph(style=style)
        font(p.add_run(block), red, bold, italic)


def base_document(title: str) -> Any:
    doc = Document()
    # Remove decorative title borders inherited from a local Word template.
    for style in doc.styles:
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.78)
    sec.left_margin = sec.right_margin = Inches(0.82)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.widow_control = True
        props = style.element.get_or_add_rPr()
        fonts = props.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            props.append(fonts)
        fonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.keep_with_next = True
    doc.core_properties.title = title
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    p = sec.footer.paragraphs[0]
    p.alignment = 2
    r = p.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    r._r.addnext(field)
    paragraph(doc, title, style="Title", bold=True)
    return doc


def evidence_text(refs: list[dict[str, Any]]) -> str:
    return "；".join(dict.fromkeys(r["locator"] for r in refs))


def stage1_doc(data: dict[str, Any]) -> Any:
    doc = base_document("审稿意见：Figure 归类与解决方案")
    meta = data["meta"]
    if text(meta.get("title")):
        paragraph(doc, meta["title"])
    comments = {c["id"]: c for c in data["comments"]}
    reviewers = {r["id"]: r for r in data["reviewers"]}
    issues = data["issues"]
    pending = sum(i["status"] == "open" for i in issues)
    paragraph(doc, f"原始意见 {len(comments)} 条；拆分问题 {len(issues)} 项；已可据实回复 {len(issues)-pending} 项；待解决 {pending} 项。")
    paragraph(doc, "每项问题只在主归属中计数，相关图不重复累计。红色标识尚待落实的处理事项。")
    if meta.get("review_sources_complete") is not True:
        paragraph(doc, "材料缺口：" + meta["source_gaps"], red=True)
    groups = sorted({i["primary_group"] for i in issues}, key=natural)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [3.05, 0.8, 1.5, 1.2]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    headers = ("主归属", "问题数", "已可据实回复", "待解决")
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
        for run in cell.paragraphs[0].runs:
            font(run, bold=True)
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    for group in groups:
        subset = [i for i in issues if i["primary_group"] == group]
        nopen = sum(i["status"] == "open" for i in subset)
        cells = table.add_row().cells
        for cell, value in zip(cells, (group, str(len(subset)), str(len(subset)-nopen), str(nopen))):
            cell.text = value
            for run in cell.paragraphs[0].runs:
                font(run)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
    for group in groups:
        paragraph(doc, group, style="Heading 1", bold=True)
        shown: set[str] = set()
        for i in [item for item in issues if item["primary_group"] == group]:
            c = comments[i["comment_id"]]
            r = reviewers[c["reviewer_id"]]
            if c["id"] not in shown:
                label = text(c.get("original_label")) or "Unnumbered comment"
                paragraph(doc, f'{r["label"]} / {label} / {c["id"]}', style="Heading 2", bold=True)
                paragraph(doc, c["original_text"], italic=True)
                shown.add(c["id"])
            red = i["status"] == "open"
            state = "待解决" if red else "已可据实回复"
            paragraph(doc, f'{i["id"]} | {state} | {i["question"]}', red=red, bold=True, style="Heading 3")
            panels = i.get("panel_refs", [])
            related = i.get("related_groups", [])
            if panels or related:
                paragraph(doc, "定位：" + "；".join(panels + ["另涉及 " + x for x in related]))
            paragraph(doc, "拟回复", bold=True, red=red)
            paragraph(doc, i["draft_response"], red=red)
            if red:
                for key, label in (("action", "解决动作"), ("inputs", "所需材料"),
                                   ("design", "实施方式"), ("output", "输出与修改位置"),
                                   ("acceptance", "核对标准"), ("fallback", "不支持或不可行时")):
                    paragraph(doc, label + "：" + i["plan"][key], red=True)
            else:
                refs = i.get("evidence", []) + i.get("resolution", {}).get("evidence", [])
                if refs:
                    paragraph(doc, "依据：" + evidence_text(refs))
    return doc


def stage2_doc(data: dict[str, Any]) -> Any:
    meta = data["meta"]
    doc = base_document("Response to the editor and reviewers")
    for key, prefix in (("title", ""), ("manuscript_id", "Manuscript ID: ")):
        if text(meta.get(key)):
            paragraph(doc, prefix + meta[key])
    paragraph(doc, meta["opening"])
    reviewers = sorted(data["reviewers"], key=lambda r: (0 if r["kind"] == "editor" else 1, r["order"]))
    for reviewer in reviewers:
        comments = sorted([c for c in data["comments"] if c["reviewer_id"] == reviewer["id"]], key=lambda c: c["order"])
        if not comments:
            continue
        paragraph(doc, reviewer["label"], style="Heading 1", bold=True)
        for c in comments:
            label = text(c.get("original_label")) or "Unnumbered comment"
            paragraph(doc, label, style="Heading 2", bold=True)
            paragraph(doc, c["original_text"], italic=True)
            paragraph(doc, "Response", style="Heading 3", bold=True)
            paragraph(doc, c["final_response"])
            kids = [i for i in data["issues"] if i["comment_id"] == c["id"]]
            refs = [ref for i in kids for ref in i.get("resolution", {}).get("changes", [])]
            if refs:
                locators = "; ".join(dict.fromkeys(ref["locator"] for ref in refs))
                paragraph(doc, "Changes in the manuscript: " + locators)
    if text(meta.get("signature")):
        paragraph(doc, meta["signature"])
    return doc


def build(data: dict[str, Any], base: Path, stage: str, output: Path) -> dict[str, Any]:
    result = validate(data, base, stage)
    output = output.resolve()
    if output.suffix.lower() != ".docx":
        raise ValidationError("output must be a .docx file")
    if output in {resolve_path(a["path"], base) for a in data["artifacts"]}:
        raise ValidationError("output would overwrite an input source")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = stage1_doc(data) if stage == "stage1" else stage2_doc(data)
    # Write only after validation; failed stage2 never creates an apparent final.
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as stream:
        tmp = Path(stream.name)
    try:
        doc.save(tmp)
        tmp.replace(output)
    finally:
        tmp.unlink(missing_ok=True)
    return {**result, "output": str(output)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("stage", choices=["stage1", "stage2", "check1", "check2"])
    parser.add_argument("ledger", type=Path)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    try:
        data = json.loads(args.ledger.read_text(encoding="utf-8-sig"))
        stage = "stage2" if args.stage in {"stage2", "check2"} else "stage1"
        if args.stage.startswith("check"):
            result = validate(data, args.ledger.resolve().parent, stage)
        else:
            if args.out is None:
                raise ValidationError("--out is required when generating Word")
            result = build(data, args.ledger.resolve().parent, stage, args.out)
        print(json.dumps(result, ensure_ascii=False))
        return 0
    except (ValidationError, OSError, ValueError, KeyError, TypeError) as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
